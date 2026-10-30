#!/usr/bin/env python3
"""Convert an industrial AI weekly Markdown report into a styled Word document."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_FONT_EAST_ASIA = "Songti SC"
BASE_FONT_LATIN = "Times New Roman"


def _set_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = BASE_FONT_LATIN
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), BASE_FONT_EAST_ASIA)
    r_fonts.set(qn("w:ascii"), BASE_FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), BASE_FONT_LATIN)
    r_fonts.set(qn("w:hint"), "eastAsia")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def _configure_style(style, size: float, bold: bool = False, color: str = "111111", before: float = 0, after: float = 6, line: float = 1.15) -> None:
    font = style.font
    font.name = BASE_FONT_LATIN
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = RGBColor.from_string(color)
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), BASE_FONT_EAST_ASIA)
    r_fonts.set(qn("w:ascii"), BASE_FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), BASE_FONT_LATIN)
    r_fonts.set(qn("w:hint"), "eastAsia")
    paragraph_format = style.paragraph_format
    paragraph_format.space_before = Pt(before)
    paragraph_format.space_after = Pt(after)
    paragraph_format.line_spacing = line


def _set_keep_with_next(paragraph, value: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if value and keep is None:
        p_pr.append(OxmlElement("w:keepNext"))
    elif not value and keep is not None:
        p_pr.remove(keep)


def _add_paragraph(document: Document, text: str, style_name: str = "Normal", align=None):
    paragraph = document.add_paragraph(style=style_name)
    run = paragraph.add_run(text)
    if style_name == "ReportTitle":
        _set_font(run, size=17, bold=True)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif style_name == "SectionHeading":
        _set_font(run, size=14, bold=True)
    elif style_name == "NewsTitle":
        _set_font(run, size=13, bold=True)
    elif style_name == "Reference":
        _set_font(run, size=9.5, bold=False, color="333333")
    else:
        _set_font(run, size=12.2, bold=False)
    if align:
        paragraph.alignment = align
    return paragraph


def _prepare_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    _configure_style(document.styles["Normal"], size=12.2, before=0, after=9, line=1.18)
    _configure_style(document.styles["Heading 1"], size=14, bold=True, before=16, after=10, line=1.12)
    _configure_style(document.styles["Heading 2"], size=13, bold=True, before=10, after=7, line=1.12)

    for style_name, base_style in [
        ("ReportTitle", "Normal"),
        ("SectionHeading", "Heading 1"),
        ("NewsTitle", "Normal"),
        ("Reference", "Normal"),
    ]:
        if style_name not in document.styles:
            document.styles.add_style(style_name, document.styles[base_style].type)

    _configure_style(document.styles["ReportTitle"], size=17, bold=True, before=0, after=18, line=1.12)
    _configure_style(document.styles["SectionHeading"], size=14, bold=True, before=20, after=12, line=1.12)
    _configure_style(document.styles["NewsTitle"], size=13, bold=True, before=14, after=9, line=1.12)
    _configure_style(document.styles["Reference"], size=9.5, bold=False, color="333333", before=0, after=3, line=1.05)
    return document


def convert(markdown_path: Path, docx_path: Path) -> None:
    document = _prepare_document()
    in_refs = False
    for raw in markdown_path.read_text(encoding="utf-8").splitlines():
        text = raw.strip()
        if not text:
            continue
        if text.startswith("# "):
            paragraph = _add_paragraph(document, text[2:].strip(), "ReportTitle")
            _set_keep_with_next(paragraph)
            continue
        if text in ("【References】", "【引用文献列表】"):
            in_refs = True
            paragraph = _add_paragraph(document, text, "SectionHeading")
            _set_keep_with_next(paragraph)
            continue
        if text.startswith("【") and text.endswith("】"):
            paragraph = _add_paragraph(document, text, "SectionHeading")
            _set_keep_with_next(paragraph)
            continue
        if in_refs or re.match(r"^\[\d+\]\s", text):
            match = re.fullmatch(r"(\[\d+\]) \[(.*?)\]\((.*?)\)", text)
            if match:
                paragraph = _add_paragraph(document, match[1] + " ", "Reference")
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(qn("r:id"), paragraph.part.relate_to(match[3], RT.HYPERLINK, is_external=True))
                run = paragraph.add_run(match[2])
                _set_font(run, size=9.5, bold=False, color="333333")
                hyperlink.append(run._r)
                paragraph._p.append(hyperlink)
            else:
                _add_paragraph(document, text, "Reference")
            continue
        if not text.endswith("。") and not text.endswith("]"):
            paragraph = _add_paragraph(document, text, "NewsTitle")
            _set_keep_with_next(paragraph)
            continue
        _add_paragraph(document, text, "Normal")
    document.save(docx_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("markdown_file", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    convert(args.markdown_file, args.out)
    print(args.out)


if __name__ == "__main__":
    main()
